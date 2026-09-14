from pathlib import Path
from html import unescape
import re
import runpy

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = '项目计划书_基于AI动态营养算法的个性化饮食推荐与健康管理系统_企业命题组修订版.docx'
SOURCE_SCRIPT = 'generate_plan_revision.py'

NAVY = '153B50'
TEAL = '0B7A75'
MINT = 'E8F5F2'
BLUE = 'EAF2F8'
GOLD = 'C58A28'
INK = '1E2933'
MUTED = '5D6B73'
LIGHT = 'F5F8FA'
LINE = 'D7E1E5'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width_twips))
    tc_w.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths_twips):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths_twips)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths_twips[min(i, len(widths_twips) - 1)])
            set_cell_margins(cell)


def set_table_borders(table, color=LINE, size='5'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_run_font(run, name='Microsoft YaHei', size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para_border_bottom(paragraph, color=TEAL, size='8', space='3'):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        p_bdr = OxmlElement('w:pBdr')
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn('w:bottom'))
    if bottom is None:
        bottom = OxmlElement('w:bottom')
        p_bdr.append(bottom)
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)


def clean_text(text):
    text = unescape(text or '')
    text = text.replace('<br/>', '\n').replace('<br />', '\n').replace('<br>', '\n')
    text = re.sub(r'<[^>]+>', '', text)
    return text


def paragraph_text(p):
    return clean_text(getattr(p, 'text', str(p)))


def add_text_to_paragraph(paragraph, text, style_name, table=False):
    lines = text.split('\n')
    for idx, line in enumerate(lines):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        if table:
            if style_name == 'TableHead':
                set_run_font(run, size=8.2, color=WHITE, bold=True)
            else:
                set_run_font(run, size=8.3, color=INK)
        elif style_name == 'CoverTitle':
            set_run_font(run, size=24, color=NAVY, bold=True)
        elif style_name == 'CoverSub':
            set_run_font(run, size=13, color=TEAL, bold=False)
        elif style_name == 'CoverMeta':
            set_run_font(run, size=10, color=MUTED)
        elif style_name == 'H1CN':
            set_run_font(run, size=16, color=NAVY, bold=True)
        elif style_name == 'H2CN':
            set_run_font(run, size=13, color=TEAL, bold=True)
        elif style_name == 'H3CN':
            set_run_font(run, size=11.5, color=INK, bold=True)
        elif style_name == 'Callout':
            set_run_font(run, size=10, color=NAVY)
        elif style_name == 'SmallCN':
            set_run_font(run, size=8.5, color=MUTED)
        else:
            set_run_font(run, size=10.5, color=INK)


def configure_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color, before, after in [
        ('Heading 1', 16, NAVY, 18, 10),
        ('Heading 2', 13, TEAL, 12, 6),
        ('Heading 3', 11.5, INK, 8, 4),
    ]:
        s = styles[name]
        s.font.name = 'Microsoft YaHei'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ['Title', 'Subtitle']:
        s = styles[name]
        s.font.name = 'Microsoft YaHei'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run('营养流 NutriFlow | 企业命题组项目计划书')
    set_run_font(run, size=8, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('基于AI动态营养算法的个性化饮食推荐与健康管理系统  |  ')
    set_run_font(run, size=8, color=MUTED)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    rfonts = OxmlElement('w:rFonts')
    rfonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    rpr.append(rfonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '16')
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    footer._p.append(fld)


def add_flow_paragraph(doc, flow):
    style_name = flow.style.name
    if style_name == 'H1CN':
        p = doc.add_paragraph(style='Heading 1')
        p.paragraph_format.keep_with_next = True
    elif style_name == 'H2CN':
        p = doc.add_paragraph(style='Heading 2')
        p.paragraph_format.keep_with_next = True
    elif style_name == 'H3CN':
        p = doc.add_paragraph(style='Heading 3')
        p.paragraph_format.keep_with_next = True
    elif style_name == 'CoverTitle':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
    elif style_name == 'CoverSub':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
    elif style_name == 'CoverMeta':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
    elif style_name == 'Callout':
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(3)
        p.paragraph_format.right_indent = Mm(3)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        set_para_border_bottom(p, color='B8D8D2', size='8', space='3')
    elif style_name == 'SmallCN':
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
    else:
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.first_line_indent = Mm(5) if style_name == 'BodyCN' else Mm(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
    add_text_to_paragraph(p, paragraph_text(flow), style_name)
    return p


def add_flow_table(doc, flow, table_idx):
    values = flow._cellvalues
    rows = len(values)
    cols = len(values[0]) if rows else 0
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    widths_pt = list(flow._colWidths or [450 / max(cols, 1)] * cols)
    widths_twips = [max(300, int(float(w) / 72 * 1440)) for w in widths_pt]
    set_table_geometry(table, widths_twips)
    set_table_borders(table)
    has_header = table_idx != 0
    for r_idx, row in enumerate(values):
        for c_idx, raw in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if has_header and r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif not has_header and c_idx == 0:
                set_cell_shading(cell, BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if hasattr(raw, 'text'):
                text = paragraph_text(raw)
                style_name = getattr(raw.style, 'name', 'TableCN')
            else:
                text = clean_text(str(raw))
                style_name = 'TableCN'
            if has_header and r_idx == 0:
                style_name = 'TableHead'
            add_text_to_paragraph(p, text, style_name, table=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    # Execute only the source's content construction, not its PDF build step.
    src = Path(SOURCE_SCRIPT).read_text(encoding='utf-8')
    namespace = {}
    exec(src.split('\ndoc = PlanDocTemplate', 1)[0], namespace)
    story = namespace['story']
    from reportlab.platypus import Paragraph, Table, PageBreak, Spacer, HRFlowable

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(19)
    section.bottom_margin = Mm(21)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    configure_styles(doc)
    add_header_footer(section)

    table_idx = 0
    for flow in story:
        if isinstance(flow, Paragraph):
            add_flow_paragraph(doc, flow)
        elif isinstance(flow, Table):
            add_flow_table(doc, flow, table_idx)
            table_idx += 1
        elif isinstance(flow, PageBreak):
            doc.add_page_break()
        elif isinstance(flow, Spacer):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(max(2, float(flow.height) * 0.7))
        elif isinstance(flow, HRFlowable):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            set_para_border_bottom(p, color=TEAL, size='8', space='1')

    doc.core_properties.title = '基于AI动态营养算法的个性化饮食推荐与健康管理系统（企业命题组修订版）'
    doc.core_properties.author = '项目团队'
    doc.core_properties.subject = '中国国际大学生创新大赛（2026）企业命题组项目计划书'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
